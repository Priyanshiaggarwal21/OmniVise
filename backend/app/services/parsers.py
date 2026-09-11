import hashlib
import mimetypes
from pathlib import Path
import re
from typing import Any, Dict, Optional, Tuple
import urllib.parse
import urllib.request
import zipfile

import fitz
from docx import Document
import pandas as pd
from pptx import Presentation


def analyze_text(text: str) -> Dict[str, Any]:
    """Smart text analyzer helper that extracts:
    - entity: finds names / corporations / counterparties
    - value: finds first significant numeric / currency value
    - date: finds ISO or formatted date strings
    - claim: extracts primary assertion or summary sentence
    """
    if not text:
        return {"entity": None, "value": None, "date": None, "claim": None}

    # 1. Entity extraction
    entity = None
    # Corporate suffix match (Inc., Corp., LLC, etc.)
    corp_pattern = r'\b([A-Z][A-Za-z0-9&.\'\-]+(?:\s+[A-Z][A-Za-z0-9&.\'\-]+)*\s+(?:Inc\.?|Corp\.?|Corporation|LLC|Ltd\.?|Co\.?|Group|Technologies|Holdings|Bank|Capital|Partners|Company))\b'
    corp_match = re.search(corp_pattern, text)
    if corp_match:
        entity = corp_match.group(1).strip()
    else:
        # Counterparty pattern (Between Acme and Beta / Party: Acme / Borrower: ...)
        party_pattern = r'(?:Between|between|Party|Client|Borrower|Lender|Vendor|Counterparty|Entity)[:\s]+([A-Z][A-Za-z0-9\s,&.\-]+?)(?=[,\n;\.]|$)'
        party_match = re.search(party_pattern, text)
        if party_match:
            candidate = party_match.group(1).strip()
            if candidate and len(candidate) < 60:
                entity = candidate
        else:
            # Capitalized noun phrases (e.g. John Doe, Acme Global)
            name_candidates = re.findall(r'\b([A-Z][a-z]+(?:\s+[A-Z][a-z]+){1,2})\b', text)
            stopwords = {
                "This Agreement", "The Company", "In Addition", "United States",
                "Executive Summary", "Table Of Contents", "Page Number", "All Rights",
                "Terms And Conditions", "North America", "South America", "Dear Sir"
            }
            for candidate in name_candidates:
                if candidate not in stopwords:
                    entity = candidate
                    break

    # 2. Value extraction (currency or significant number)
    value: Optional[float] = None
    curr_match = re.search(r'[\$€£]\s*([0-9]{1,3}(?:,[0-9]{3})+(?:\.[0-9]+)?|[0-9]+(?:\.[0-9]+)?)', text)
    if curr_match:
        try:
            value = float(curr_match.group(1).replace(",", ""))
        except ValueError:
            value = None

    if value is None:
        iso_curr_match = re.search(r'\b(?:USD|EUR|GBP|INR)\s*([0-9]{1,3}(?:,[0-9]{3})+(?:\.[0-9]+)?|[0-9]+(?:\.[0-9]+)?)\b', text)
        if iso_curr_match:
            try:
                value = float(iso_curr_match.group(1).replace(",", ""))
            except ValueError:
                value = None

    if value is None:
        num_match = re.search(r'\b([0-9]{1,3}(?:,[0-9]{3})+(?:\.[0-9]+)?)\b', text)
        if num_match:
            try:
                value = float(num_match.group(1).replace(",", ""))
            except ValueError:
                value = None

    if value is None:
        float_match = re.search(r'\b([0-9]+\.[0-9]+)\b', text)
        if float_match:
            try:
                value = float(float_match.group(1))
            except ValueError:
                value = None

    # 3. Date extraction
    extracted_date: Optional[str] = None
    iso_date_match = re.search(r'\b(\d{4}-\d{2}-\d{2})\b', text)
    if iso_date_match:
        extracted_date = iso_date_match.group(1)
    else:
        named_date_match = re.search(
            r'\b((?:January|February|March|April|May|June|July|August|September|October|November|December|Jan|Feb|Mar|Apr|Jun|Jul|Aug|Sep|Oct|Nov|Dec)[a-z]*\.?\s+\d{1,2}(?:st|nd|rd|th)?,?\s+\d{4})\b',
            text,
            re.IGNORECASE
        )
        if named_date_match:
            extracted_date = named_date_match.group(1)
        else:
            slash_date_match = re.search(r'\b(\d{1,2}[/-]\d{1,2}[/-]\d{2,4})\b', text)
            if slash_date_match:
                extracted_date = slash_date_match.group(1)

    # 4. Claim extraction
    claim: Optional[str] = None
    sentences = [s.strip() for s in re.split(r'[\r\n]+|(?<=[.!?])\s+', text) if s.strip()]
    claim_keywords = [
        "revenue", "increased", "decreased", "totaled", "reported", "net income",
        "sales", "earnings", "growth", "agreed", "shall", "warrants", "asserts",
        "claims", "disclosed", "guidance", "profit", "contract", "liability", "assets"
    ]
    for s in sentences:
        if any(k in s.lower() for k in claim_keywords) and 15 <= len(s) <= 300:
            claim = s
            break

    if not claim and sentences:
        for s in sentences:
            if len(s) >= 15:
                claim = s[:300]
                break

    return {
        "entity": entity,
        "value": value,
        "date": extracted_date,
        "claim": claim
    }


def pdf_parser(file_path: Path, filename: str) -> dict:
    """Extracts text, page count, and metadata using pymupdf."""
    doc = fitz.open(str(file_path))
    page_count = len(doc)
    text_chunks = []
    for page in doc:
        text_chunks.append(page.get_text())
    metadata = doc.metadata or {}
    doc.close()

    full_text = "\n".join(text_chunks).strip()
    analysis = analyze_text(full_text)
    claim = analysis["claim"] or f"PDF document {filename} with {page_count} pages processed."

    return {
        "source": filename,
        "modality": "pdf",
        "claim": claim,
        "entity": analysis["entity"],
        "value": analysis["value"],
        "date": analysis["date"],
        "location": f"Page 1 of {page_count}" if page_count > 0 else "Page 1",
        "raw_text_preview": full_text[:500] if full_text else "",
        "metadata": {
            "page_count": page_count,
            "title": metadata.get("title", ""),
            "author": metadata.get("author", ""),
            "subject": metadata.get("subject", ""),
            "filename": filename,
        }
    }


def docx_parser(file_path: Path, filename: str) -> dict:
    """Extracts paragraphs and table text using python-docx."""
    doc = Document(str(file_path))
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    table_texts = []
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                table_texts.append(" | ".join(cells))

    full_text = "\n".join(paragraphs + table_texts).strip()
    analysis = analyze_text(full_text)
    claim = analysis["claim"] or f"Word document {filename} processed with {len(paragraphs)} paragraphs."

    return {
        "source": filename,
        "modality": "docx",
        "claim": claim,
        "entity": analysis["entity"],
        "value": analysis["value"],
        "date": analysis["date"],
        "location": f"Paragraph 1 ({len(paragraphs)} paragraphs, {len(doc.tables)} tables)",
        "raw_text_preview": full_text[:500] if full_text else "",
        "metadata": {
            "paragraph_count": len(paragraphs),
            "table_count": len(doc.tables),
            "filename": filename,
        }
    }


def pptx_parser(file_path: Path, filename: str) -> dict:
    """Extracts slide shapes text using python-pptx."""
    prs = Presentation(str(file_path))
    slide_texts = []
    for i, slide in enumerate(prs.slides):
        slide_parts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for p in shape.text_frame.paragraphs:
                    t = p.text.strip()
                    if t:
                        slide_parts.append(t)
        if slide_parts:
            slide_texts.append(f"Slide {i + 1}: " + " ".join(slide_parts))

    full_text = "\n".join(slide_texts).strip()
    analysis = analyze_text(full_text)
    slide_count = len(prs.slides)
    claim = analysis["claim"] or f"Presentation {filename} with {slide_count} slides processed."

    return {
        "source": filename,
        "modality": "pptx",
        "claim": claim,
        "entity": analysis["entity"],
        "value": analysis["value"],
        "date": analysis["date"],
        "location": f"Slide 1 of {slide_count}" if slide_count > 0 else "Slide 1",
        "raw_text_preview": full_text[:500] if full_text else "",
        "metadata": {
            "slide_count": slide_count,
            "filename": filename,
        }
    }


def excel_parser(file_path: Path, filename: str) -> dict:
    """Extracts sheet names, rows, columns, numerical summaries using pandas / openpyxl."""
    is_csv = filename.lower().endswith(".csv")
    sheets = {}
    if is_csv:
        try:
            df = pd.read_csv(str(file_path))
            sheets = {"Sheet1": df}
        except Exception:
            sheets = {"Sheet1": pd.DataFrame()}
    else:
        try:
            xls = pd.ExcelFile(str(file_path))
            sheets = {name: pd.read_excel(xls, sheet_name=name) for name in xls.sheet_names}
        except Exception:
            sheets = {"Sheet1": pd.DataFrame()}

    total_rows = 0
    total_cols = 0
    sheet_names = list(sheets.keys())
    numerical_summaries = {}
    first_numeric_val = None
    full_text_parts = []

    for sname, df in sheets.items():
        rows, cols = df.shape
        total_rows += rows
        total_cols += cols
        num_cols = df.select_dtypes(include=["number"]).columns.tolist()
        num_summary = {}
        for c in num_cols:
            col_series = df[c].dropna()
            if not col_series.empty:
                col_sum = float(col_series.sum())
                num_summary[str(c)] = {
                    "sum": col_sum,
                    "mean": float(col_series.mean()),
                    "min": float(col_series.min()),
                    "max": float(col_series.max()),
                }
                if first_numeric_val is None:
                    first_numeric_val = float(col_series.iloc[0])
        numerical_summaries[sname] = num_summary
        head_str = df.head(5).to_string()
        full_text_parts.append(f"Sheet {sname} ({rows}x{cols}):\n{head_str}")

    combined_text = "\n\n".join(full_text_parts)
    analysis = analyze_text(combined_text)
    if first_numeric_val is not None and analysis["value"] is None:
        analysis["value"] = first_numeric_val

    first_sheet = sheet_names[0] if sheet_names else "Sheet1"
    claim = analysis["claim"] or f"Spreadsheet {filename} with {len(sheet_names)} sheet(s) and {total_rows} rows."

    return {
        "source": filename,
        "modality": "excel",
        "claim": claim,
        "entity": analysis["entity"],
        "value": analysis["value"],
        "date": analysis["date"],
        "location": f"Sheet: {first_sheet}, Rows: {total_rows}, Cols: {total_cols}",
        "raw_text_preview": combined_text[:500] if combined_text else "",
        "metadata": {
            "sheet_names": sheet_names,
            "total_rows": total_rows,
            "total_columns": total_cols,
            "numerical_summaries": numerical_summaries,
            "filename": filename,
        }
    }


def audio_parser(file_path: Path, filename: str) -> dict:
    """Returns stub metadata for audio with modality='audio',
    claim='Audio stream registered for diarization and Whisper transcription', location='00:00:00'.
    """
    return {
        "source": filename,
        "modality": "audio",
        "claim": "Audio stream registered for diarization and Whisper transcription",
        "entity": None,
        "value": None,
        "date": None,
        "location": "00:00:00",
        "raw_text_preview": None,
        "metadata": {
            "filename": filename,
            "pipeline": "whisper_diarization_pipeline",
            "status": "registered",
        }
    }


def video_parser(file_path: Path, filename: str) -> dict:
    """Returns stub metadata for video with modality='video',
    claim='Video stream registered for frame OCR and transcript extraction', location='00:00:00'.
    """
    return {
        "source": filename,
        "modality": "video",
        "claim": "Video stream registered for frame OCR and transcript extraction",
        "entity": None,
        "value": None,
        "date": None,
        "location": "00:00:00",
        "raw_text_preview": None,
        "metadata": {
            "filename": filename,
            "pipeline": "video_frame_ocr_transcript_pipeline",
            "status": "registered",
        }
    }


def web_parser(url: str, temp_path: Optional[Path] = None) -> dict:
    """Fetches URL or parses domain, extracts text/title, returns modality='web'."""
    parsed = urllib.parse.urlparse(url)
    domain = parsed.netloc or url
    path = parsed.path
    title = ""
    body_text = ""

    try:
        req = urllib.request.Request(
            url,
            headers={"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) OmniViseEvidenceEngine/1.0"}
        )
        with urllib.request.urlopen(req, timeout=3.0) as resp:
            content_bytes = resp.read(65536)
            html = content_bytes.decode("utf-8", errors="ignore")
            title_match = re.search(r'<title[^>]*>(.*?)</title>', html, re.IGNORECASE | re.DOTALL)
            if title_match:
                title = title_match.group(1).strip()
            body_text = re.sub(r'<[^>]+>', ' ', html)
            body_text = re.sub(r'\s+', ' ', body_text).strip()
    except Exception:
        pass

    if not title:
        clean_path = path.strip("/").replace("/", " - ").replace("_", " ").replace("-", " ")
        title = f"{domain}: {clean_path}" if clean_path else domain

    combined = f"{title}\n{body_text}".strip()
    analysis = analyze_text(combined)
    claim = analysis["claim"] or f"Web evidence extracted from {domain} ({title})"

    return {
        "source": url,
        "modality": "web",
        "claim": claim,
        "entity": analysis["entity"] or domain,
        "value": analysis["value"],
        "date": analysis["date"],
        "location": url,
        "raw_text_preview": body_text[:500] if body_text else title,
        "metadata": {
            "url": url,
            "domain": domain,
            "title": title,
        }
    }


def zip_parser(file_path: Path, filename: str) -> dict:
    """Inspects archive with zipfile.ZipFile, lists files, returns modality='zip'."""
    namelist = []
    total_size = 0
    sample_text = ""

    try:
        with zipfile.ZipFile(str(file_path), "r") as zf:
            namelist = zf.namelist()
            total_size = sum(info.file_size for info in zf.infolist())
            for name in namelist:
                if name.lower().endswith((".txt", ".csv", ".json", ".md")) and not name.startswith("__MACOSX"):
                    try:
                        sample_text = zf.read(name)[:1000].decode("utf-8", errors="ignore")
                        break
                    except Exception:
                        pass
    except Exception:
        namelist = []

    analysis = analyze_text(sample_text) if sample_text else {"entity": None, "value": None, "date": None, "claim": None}
    file_count = len(namelist)
    claim = analysis["claim"] or f"Archive {filename} containing {file_count} file(s): {', '.join(namelist[:5])}"

    return {
        "source": filename,
        "modality": "zip",
        "claim": claim,
        "entity": analysis["entity"],
        "value": analysis["value"],
        "date": analysis["date"],
        "location": f"{file_count} files in archive",
        "raw_text_preview": sample_text[:500] if sample_text else f"Archive containing: {', '.join(namelist[:10])}",
        "metadata": {
            "file_count": file_count,
            "file_list": namelist[:50],
            "total_uncompressed_size_bytes": total_size,
            "filename": filename,
        }
    }


def text_parser(file_path: Path, filename: str) -> dict:
    """Extracts text content from plain text files."""
    try:
        content = file_path.read_text(encoding="utf-8", errors="ignore")
    except Exception:
        content = ""

    analysis = analyze_text(content)
    claim = analysis["claim"] or f"Text document {filename} processed."

    return {
        "source": filename,
        "modality": "text",
        "claim": claim,
        "entity": analysis["entity"],
        "value": analysis["value"],
        "date": analysis["date"],
        "location": "Line 1",
        "raw_text_preview": content[:500] if content else "",
        "metadata": {
            "char_count": len(content),
            "filename": filename,
        }
    }


def parse_evidence(
    file_path: Optional[Path],
    filename: Optional[str],
    url: Optional[str],
    mime_type: Optional[str]
) -> Tuple[dict, str]:
    """Dispatcher that determines parser and returns (parsed_data, parser_name)."""
    if url:
        return web_parser(url, file_path), "web_parser"

    if not file_path:
        raise ValueError("Either file_path or url must be provided.")

    fn = (filename or file_path.name).lower()
    mime = (mime_type or "").lower()

    if fn.endswith(".pdf") or "pdf" in mime:
        return pdf_parser(file_path, filename or file_path.name), "pdf_parser"
    elif fn.endswith((".docx", ".doc")) or "wordprocessingml" in mime or "msword" in mime:
        return docx_parser(file_path, filename or file_path.name), "docx_parser"
    elif fn.endswith((".pptx", ".ppt")) or "presentationml" in mime or "powerpoint" in mime:
        return pptx_parser(file_path, filename or file_path.name), "pptx_parser"
    elif fn.endswith((".xlsx", ".xls", ".csv")) or "spreadsheetml" in mime or "excel" in mime or "csv" in mime:
        return excel_parser(file_path, filename or file_path.name), "excel_parser"
    elif fn.endswith((".mp3", ".wav", ".m4a", ".ogg", ".flac", ".aac")) or mime.startswith("audio/"):
        return audio_parser(file_path, filename or file_path.name), "audio_parser"
    elif fn.endswith((".mp4", ".mov", ".avi", ".mkv", ".webm", ".m4v")) or mime.startswith("video/"):
        return video_parser(file_path, filename or file_path.name), "video_parser"
    elif fn.endswith((".zip", ".tar", ".gz", ".7z")) or "zip" in mime or "compressed" in mime or "archive" in mime:
        return zip_parser(file_path, filename or file_path.name), "zip_parser"
    else:
        return text_parser(file_path, filename or file_path.name), "text_parser"
