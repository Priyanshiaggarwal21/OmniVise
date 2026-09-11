import io
import unittest
import zipfile

from docx import Document
import fitz
from pptx import Presentation
from starlette.testclient import TestClient

from app.main import app


class TestEvidenceUploadEndpoint(unittest.TestCase):
    def setUp(self):
        self.client = TestClient(app)
        self.required_fields = {
            "source",
            "modality",
            "claim",
            "entity",
            "value",
            "date",
            "version",
            "location",
            "provenance",
        }
        self.required_provenance_fields = {
            "file_name",
            "file_hash",
            "temp_path",
            "parser_used",
            "file_size_bytes",
            "mime_type",
            "ingested_at",
            "storage_type",
        }

    def _assert_schema(self, data: dict):
        for field in self.required_fields:
            self.assertIn(field, data, f"Missing required evidence field: {field}")
        prov = data["provenance"]
        self.assertIsInstance(prov, dict)
        for pfield in self.required_provenance_fields:
            self.assertIn(pfield, prov, f"Missing required provenance field: {pfield}")

    def test_upload_text_file(self):
        content = b"Acme Corp reported $1500000 revenue on 2024-05-10."
        response = self.client.post(
            "/evidence/upload",
            files={"file": ("report.txt", content, "text/plain")},
        )
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self._assert_schema(data)
        self.assertEqual(data["modality"], "text")
        self.assertEqual(data["entity"], "Acme Corp")
        self.assertEqual(data["value"], 1500000.0)
        self.assertEqual(data["date"], "2024-05-10")
        self.assertEqual(data["provenance"]["parser_used"], "text_parser")

    def test_upload_csv_file(self):
        csv_content = (
            b"Entity,Quarter,Revenue,Date\n"
            b"Acme Corp,Q1,500000.00,2024-03-31\n"
            b"Acme Corp,Q2,750000.00,2024-06-30\n"
        )
        response = self.client.post(
            "/api/v1/evidence/upload",
            files={"file": ("financials.csv", csv_content, "text/csv")},
        )
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self._assert_schema(data)
        self.assertEqual(data["modality"], "excel")
        self.assertEqual(data["provenance"]["parser_used"], "excel_parser")
        self.assertIn("Sheet1", data["metadata"]["sheet_names"])

    def test_upload_pdf_file(self):
        doc = fitz.open()
        page = doc.new_page()
        page.insert_text((50, 72), "Alphabet Inc reported $75000.50 net income on 2024-04-25.")
        pdf_bytes = doc.tobytes()
        doc.close()

        response = self.client.post(
            "/evidence/upload",
            files={"file": ("q1_filing.pdf", pdf_bytes, "application/pdf")},
        )
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self._assert_schema(data)
        self.assertEqual(data["modality"], "pdf")
        self.assertEqual(data["entity"], "Alphabet Inc")
        self.assertEqual(data["value"], 75000.5)
        self.assertEqual(data["date"], "2024-04-25")
        self.assertEqual(data["provenance"]["parser_used"], "pdf_parser")
        self.assertIn("Page 1", data["location"])

    def test_upload_docx_file(self):
        doc = Document()
        doc.add_paragraph("Microsoft Corp disclosed $5000000 in quarterly revenue on 2024-02-15.")
        buf = io.BytesIO()
        doc.save(buf)
        docx_bytes = buf.getvalue()

        response = self.client.post(
            "/evidence/upload",
            files={
                "file": (
                    "contract.docx",
                    docx_bytes,
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
        )
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self._assert_schema(data)
        self.assertEqual(data["modality"], "docx")
        self.assertEqual(data["entity"], "Microsoft Corp")
        self.assertEqual(data["value"], 5000000.0)
        self.assertEqual(data["date"], "2024-02-15")
        self.assertEqual(data["provenance"]["parser_used"], "docx_parser")

    def test_upload_pptx_file(self):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        slide.shapes.title.text = "Tesla Inc reported $8200000 in battery sales on 2024-03-01."
        buf = io.BytesIO()
        prs.save(buf)
        pptx_bytes = buf.getvalue()

        response = self.client.post(
            "/evidence/upload",
            files={
                "file": (
                    "earnings_deck.pptx",
                    pptx_bytes,
                    "application/vnd.openxmlformats-officedocument.presentationml.presentation",
                )
            },
        )
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self._assert_schema(data)
        self.assertEqual(data["modality"], "pptx")
        self.assertEqual(data["entity"], "Tesla Inc")
        self.assertEqual(data["value"], 8200000.0)
        self.assertEqual(data["date"], "2024-03-01")
        self.assertEqual(data["provenance"]["parser_used"], "pptx_parser")

    def test_upload_audio_stub(self):
        response = self.client.post(
            "/evidence/upload",
            files={"file": ("earnings_call.mp3", b"\x00\x01\x02\x03", "audio/mpeg")},
        )
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self._assert_schema(data)
        self.assertEqual(data["modality"], "audio")
        self.assertEqual(
            data["claim"],
            "Audio stream registered for diarization and Whisper transcription",
        )
        self.assertEqual(data["location"], "00:00:00")
        self.assertEqual(data["provenance"]["parser_used"], "audio_parser")

    def test_upload_video_stub(self):
        response = self.client.post(
            "/evidence/upload",
            files={"file": ("press_conference.mp4", b"\x00\x00\x00\x18ftypmp42", "video/mp4")},
        )
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self._assert_schema(data)
        self.assertEqual(data["modality"], "video")
        self.assertEqual(
            data["claim"],
            "Video stream registered for frame OCR and transcript extraction",
        )
        self.assertEqual(data["location"], "00:00:00")
        self.assertEqual(data["provenance"]["parser_used"], "video_parser")

    def test_upload_zip_archive(self):
        zbuf = io.BytesIO()
        with zipfile.ZipFile(zbuf, "w") as zf:
            zf.writestr("notes.txt", "Internal audit notes for FY2024.")
            zf.writestr("readme.md", "README for bundle.")
        zip_bytes = zbuf.getvalue()

        response = self.client.post(
            "/evidence/upload",
            files={"file": ("evidence_bundle.zip", zip_bytes, "application/zip")},
        )
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self._assert_schema(data)
        self.assertEqual(data["modality"], "zip")
        self.assertEqual(data["provenance"]["parser_used"], "zip_parser")
        self.assertEqual(data["metadata"]["file_count"], 2)

    def test_upload_url_via_form_data(self):
        response = self.client.post(
            "/evidence/upload",
            data={"url": "https://example.com/sec/10q"},
        )
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self._assert_schema(data)
        self.assertEqual(data["modality"], "web")
        self.assertEqual(data["source"], "https://example.com/sec/10q")
        self.assertEqual(data["provenance"]["parser_used"], "web_parser")

    def test_upload_url_via_json_body(self):
        response = self.client.post(
            "/api/v1/evidence/upload",
            json={"url": "https://example.com/sec/10q"},
        )
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self._assert_schema(data)
        self.assertEqual(data["modality"], "web")
        self.assertEqual(data["source"], "https://example.com/sec/10q")
        self.assertEqual(data["provenance"]["parser_used"], "web_parser")

    def test_missing_file_and_url_raises_400(self):
        # Empty multipart form
        response1 = self.client.post("/evidence/upload", data={})
        self.assertEqual(response1.status_code, 400)
        self.assertIn("Either 'file' or 'url' must be provided", response1.json()["detail"])

        # Empty JSON body
        response2 = self.client.post("/evidence/upload", json={})
        self.assertEqual(response2.status_code, 400)
        self.assertIn("Either 'file' or 'url' must be provided", response2.json()["detail"])

    def test_schema_verification(self):
        response = self.client.post(
            "/evidence/upload",
            data={"url": "https://example.com/sec/10q"},
        )
        self.assertEqual(response.status_code, 200)
        data = response.json()

        # Check all 9 specific fields from specification
        spec_fields = ["source", "modality", "claim", "entity", "value", "date", "version", "location", "provenance"]
        for sf in spec_fields:
            self.assertIn(sf, data, f"Required spec field '{sf}' missing from response")

        # Provenance fields verification
        self.assertIn("file_name", data["provenance"])
        self.assertIn("file_hash", data["provenance"])
        self.assertIn("parser_used", data["provenance"])
        self.assertIn("file_size_bytes", data["provenance"])
        self.assertIn("mime_type", data["provenance"])
        self.assertIn("ingested_at", data["provenance"])
        self.assertEqual(data["provenance"]["storage_type"], "ephemeral_temp")


if __name__ == "__main__":
    unittest.main()
